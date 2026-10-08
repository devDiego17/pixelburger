from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

OUT = r"C:\Users\diego\OneDrive\Documentos\PixelBurger\Mision Pixel - Sistema de XP.xlsx"
NJ = 500    # filas de jugadores
NM = 3000   # filas de movimientos
JL = NJ + 1
ML = NM + 1

F = "Arial"
H_FILL = PatternFill("solid", fgColor="15131F")
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
IN_FONT = Font(name=F, color="0000FF", size=10)
CALC_FONT = Font(name=F, color="000000", size=10)
BASE = Font(name=F, size=10)
EX_FILL = PatternFill("solid", fgColor="FFF6D5")
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE = Font(name=F, bold=True, size=14, color="C2185B")
SUB = Font(name=F, bold=True, size=11)
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=thin)

wb = Workbook()

def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = H_FONT
        c.fill = H_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 32

# ---------------- Cómo usar ----------------
ws = wb.active
ws.title = "Cómo usar"
ws.column_dimensions["A"].width = 120
lines = [
    ("MISIÓN PIXEL · SISTEMA DE XP", TITLE),
    ("Reto por puntos de Pixel Burger. Registros desde el 20 de octubre, compras con XP desde el 1 de noviembre, cierre el 18 de noviembre a las 23:59 y premiación el 19 de noviembre en el local.", BASE),
    ("", BASE),
    ("CÓMO SE USA ESTE ARCHIVO", SUB),
    ("Texto azul = celdas que se llenan a mano. Celdas grises = fórmulas: no se tocan. Fila amarilla = ejemplo: bórrala antes de empezar.", BASE),
    ("1. Jugadores: una fila por persona registrada (se copia del formulario). El ID es #0001, #0002… en orden y nunca se reutiliza.", BASE),
    ("2. Movimientos: cada vez que alguien gana XP se agrega UNA fila (fecha y hora, ID, acción, cantidad, quién lo registró). La XP se calcula sola.", BASE),
    ("3. Ranking: se actualiza solo. Muestra solo el alias público, nunca el nombre ni el teléfono.", BASE),
    ("4. Tabla XP y Logros: aquí se cambian los puntos y las recompensas. Si cambias un valor, todo se recalcula.", BASE),
    ("", BASE),
    ("EN LA CAJA (cada pedido)", SUB),
    ("1. Preguntar: «¿Tienes número de jugador?». Si no lo recuerda, buscarlo por WhatsApp en la hoja Jugadores.", BASE),
    ("2. Agregar en Movimientos: «Pedido» (cantidad 1) + «Hamburguesa adicional» (cantidad = hamburguesas extra, máximo 4) + «Bonus Montañera» (cantidad = Montañeras).", BASE),
    ("3. Revisar en Jugadores la columna «¿Entregar recompensa?». Si dice Sí, entregarla y actualizar «Nivel entregado».", BASE),
    ("4. Enviar el WhatsApp de XP sumada (plantillas en la hoja Mensajes).", BASE),
    ("", BASE),
    ("REGLAS CLAVE (resumen; las bases completas deben revisarse con un abogado antes de publicarse)", SUB),
    ("• Solo mayores de 18 años, con autorización de tratamiento de datos (Ley 1581 de 2012). Un jugador = un número de WhatsApp.", BASE),
    ("• La XP de compra solo cuenta si el jugador da su número al pagar. No se suma XP retroactiva.", BASE),
    ("• Máximo 1 pedido con XP por día. Los límites de cada acción están en la hoja Tabla XP.", BASE),
    ("• Desempate: gana quien llegó primero a ese puntaje (fecha de su último movimiento). No hay azar.", BASE),
    ("• Empleados, socios y sus familiares no compiten por los premios del ranking.", BASE),
    ("• Pixel Burger puede anular XP obtenida con trampa (se registra como «Ajuste del administrador» con XP negativa).", BASE),
    ("• Control cruzado: cada lunes el socio que no registró los movimientos revisa la hoja Movimientos.", BASE),
    ("", BASE),
    ("PREMIOS DEL RANKING (19 de noviembre)", SUB),
    ("1.er puesto: GTA VI (PS5 o Xbox Series, a elección del ganador) · 2.º y 3.º: 2 hamburguesas a elección · 4.º al 10.º: 1 hamburguesa a elección.", BASE),
    ("Además, cada jugador desbloquea recompensas por nivel (hoja Logros). La XP no se gasta: los niveles se acumulan.", BASE),
]
for i, (t, f) in enumerate(lines, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")

# ---------------- Tabla XP ----------------
tx = wb.create_sheet("Tabla XP")
header(tx, 1, ["Acción", "XP", "Límite", "Cómo se verifica", "Por qué"], [42, 8, 30, 42, 46])
acciones = [
    ("Registro", 100, "1 vez por jugador", "Formulario completo, mayor de 18, autoriza datos", "Entrada al juego"),
    ("Pedido", 300, "Máximo 1 pedido con XP al día", "El cajero anota el ID al cobrar", "Premia la frecuencia de compra"),
    ("Hamburguesa adicional en el pedido", 100, "Máximo 4 por pedido", "Mismo ticket del pedido", "Sube el ticket sin que el ranking se compre con un solo pedido grande"),
    ("Bonus Montañera", 50, "Por cada Montañera del pedido", "Mismo ticket del pedido", "Es la hamburguesa que más margen deja"),
    ("Foto con Byte publicada", 200, "Máximo 1 por semana", "Publicación pública que etiqueta a Pixel Burger", "Contenido y alcance en redes"),
    ("Referido con primer pedido", 400, "Máximo 5 referidos por jugador", "El referido escribió este ID al registrarse y ya hizo su primer pedido", "Clientes nuevos que sí compran"),
    ("Diseño de caja #CajaPixel", 300, "Máximo 1 por semana", "Foto pública con #CajaPixel o enviada por WhatsApp", "Conecta con «No tenemos imagen»"),
    ("Código del cartel secreto", 150, "Cada código, 1 vez por jugador", "El jugador envía el código por WhatsApp", "Hace que la gente busque los carteles"),
    ("Le ganó a Byte", 100, "Máximo 1 al día", "Lo confirma el acompañante de Byte", "Interacción en la calle"),
    ("Ajuste del administrador", 1, "Solo el administrador", "La nota explica el motivo", "Corregir errores o anular trampas: la XP = la cantidad (puede ser negativa)"),
    ("Recompensa entregada", 0, "Informativo", "Nota: qué se entregó", "Deja rastro de los premios entregados"),
]
for r, row in enumerate(acciones, start=2):
    for c, v in enumerate(row, start=1):
        cell = tx.cell(row=r, column=c, value=v)
        cell.font = IN_FONT if c == 2 else BASE
        cell.alignment = Alignment(wrap_text=True, vertical="top")
tx.freeze_panes = "A2"
ACC_RANGE = f"'Tabla XP'!$A$2:$A${1 + len(acciones)}"
XP_RANGE = f"'Tabla XP'!$B$2:$B${1 + len(acciones)}"

# ---------------- Logros ----------------
lg = wb.create_sheet("Logros")
header(lg, 1, ["Nivel", "Nombre", "XP mínima", "Recompensa (1 vez por jugador)", "Costo aprox. para Pixel Burger", "Equivale más o menos a"], [8, 14, 12, 52, 30, 40])
logros = [
    (0, "Sin nivel", 0, "—", "—", "—"),
    (1, "Jugador", 100, "Power-up: salsa extra gratis en el primer pedido", "≈ $200–400", "Registrarse"),
    (2, "Veterano", 1000, "Mejora gratis: Montañera o Cheese al precio de la Clásica", "≈ $1.300–2.200 (la diferencia de costo)", "Registro + 3 pedidos"),
    (3, "Élite", 2000, "Una Clásica gratis", "≈ $6.900 (ingredientes + empaque)", "~6 pedidos, o 4 pedidos + 1 referido"),
    (4, "Leyenda", 4000, "Dos hamburguesas Clásicas gratis", "≈ $13.800", "~10–12 acciones"),
]
for r, row in enumerate(logros, start=2):
    for c, v in enumerate(row, start=1):
        cell = lg.cell(row=r, column=c, value=v)
        cell.font = IN_FONT if c in (3, 4) else BASE
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if c == 3:
            cell.number_format = "#,##0"
n_log = len(logros)
lg.cell(row=n_log + 3, column=1, value="La XP mínima debe ir de menor a mayor. Los costos son estimados con los costos del plan recomendado (octubre 2026); el costo de cada recompensa es ~2–6% del margen que dejan los pedidos necesarios para alcanzarla.").font = Font(name=F, italic=True, size=9)
LOG_LVL = f"Logros!$A$2:$A${1 + n_log}"
LOG_XP = f"Logros!$C$2:$C${1 + n_log}"
LOG_REW = f"Logros!$D$2:$D${1 + n_log}"

# ---------------- Jugadores ----------------
jg = wb.create_sheet("Jugadores")
jcols = ["ID", "Nombre", "Alias público (ranking)", "WhatsApp", "Instagram", "Barrio", "Fecha de registro",
         "¿Mayor de 18?", "¿Autoriza datos?", "Referido por (ID)", "¿Empleado o familiar?",
         "XP total", "Último movimiento", "¿Compite?", "Puesto", "Nivel alcanzado", "Nivel entregado",
         "¿Entregar recompensa?", "Clave orden (no tocar)"]
jw = [8, 22, 20, 14, 18, 16, 16, 10, 10, 12, 11, 10, 17, 10, 8, 10, 10, 34, 12]
header(jg, 1, jcols, jw)
jg.freeze_panes = "C2"
example = ["#0001", "Ejemplo (borrar)", "PlayerUno", "3001234567", "@ejemplo", "Robledo", datetime(2026, 10, 20, 18, 30),
           "Sí", "Sí", "", "No"]
for r in range(2, JL + 1):
    # entradas
    if r == 2:
        for c, v in enumerate(example, start=1):
            cell = jg.cell(row=r, column=c, value=v)
    for c in range(1, 12):
        cell = jg.cell(row=r, column=c)
        cell.font = IN_FONT
        if c == 7:
            cell.number_format = "yyyy-mm-dd hh:mm"
        if c == 4:
            cell.number_format = "@"
    jg.cell(row=r, column=17).font = IN_FONT  # nivel entregado
    if r == 2:
        jg.cell(row=r, column=17, value=1)
    # fórmulas
    f = {
        12: f'=IF(A{r}="","",SUMIF(Movimientos!$B$2:$B${ML},A{r},Movimientos!$F$2:$F${ML}))',
        13: f'=IF(A{r}="","",IF(COUNTIF(Movimientos!$B$2:$B${ML},A{r})=0,"",_xlfn.MAXIFS(Movimientos!$A$2:$A${ML},Movimientos!$B$2:$B${ML},A{r})))',
        14: f'=IF(A{r}="","",IF(AND(H{r}="Sí",I{r}="Sí",K{r}<>"Sí"),"Sí","No"))',
        19: f'=IF(OR(A{r}="",N{r}<>"Sí",N(L{r})<=0),"",(N(L{r})*100000-MAX(0,ROUND((N(M{r})-DATE(2026,10,1))*1440,0)))*1000-ROW())',
        15: f'=IF(S{r}="","",1+COUNTIFS($N$2:$N${JL},"Sí",$S$2:$S${JL},">"&S{r}))',
        16: f'=IF(A{r}="","",INDEX({LOG_LVL},MATCH(N(L{r}),{LOG_XP},1)))',
        18: f'=IF(OR(A{r}="",P{r}=""),"",IF(P{r}>N(Q{r}),"Sí: entregar nivel "&(N(Q{r})+1)&IF(P{r}>N(Q{r})+1," al "&P{r},"")&" (ver Logros)",""))',
    }
    for c, formula in f.items():
        cell = jg.cell(row=r, column=c, value=formula)
        cell.font = CALC_FONT
        cell.fill = CALC_FILL
        if c == 12:
            cell.number_format = "#,##0"
        if c == 13:
            cell.number_format = "yyyy-mm-dd hh:mm"
        if c == 19:
            cell.number_format = "0"
for c in range(1, 12):
    jg.cell(row=2, column=c).fill = EX_FILL
jg.column_dimensions["S"].hidden = True
jg.cell(row=1, column=17).comment = Comment("Lo llena quien entrega la recompensa: el último nivel cuya recompensa ya se entregó (0 si ninguna).", "Pixel Burger")
jg.cell(row=1, column=15).comment = Comment("Desempate: con igual XP gana quien llegó primero (último movimiento más antiguo). Solo compiten mayores de 18 con autorización de datos y que no sean empleados o familiares.", "Pixel Burger")

dv_si = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
jg.add_data_validation(dv_si)
for col in ("H", "I", "K"):
    dv_si.add(f"{col}2:{col}{JL}")
dv_lvl = DataValidation(type="whole", operator="between", formula1="0", formula2=str(n_log - 1), allow_blank=True)
jg.add_data_validation(dv_lvl)
dv_lvl.add(f"Q2:Q{JL}")
jg.conditional_formatting.add(f"R2:R{JL}", FormulaRule(formula=[f'LEFT(R2,2)="Sí"'], fill=PatternFill("solid", fgColor="FFE08A"), font=Font(name=F, bold=True)))
jg.conditional_formatting.add(f"A2:A{JL}", FormulaRule(formula=[f'AND(A2<>"",COUNTIF($A$2:$A${JL},A2)>1)'], fill=PatternFill("solid", fgColor="F8C9C9")))
jg.conditional_formatting.add(f"D2:D{JL}", FormulaRule(formula=[f'AND(D2<>"",COUNTIF($D$2:$D${JL},D2)>1)'], fill=PatternFill("solid", fgColor="F8C9C9")))

# ---------------- Movimientos ----------------
mv = wb.create_sheet("Movimientos")
header(mv, 1, ["Fecha y hora", "ID jugador", "Alias (automático)", "Acción", "Cantidad", "XP", "Registró", "Evidencia / nota"],
       [17, 11, 22, 36, 10, 9, 14, 44])
mv.freeze_panes = "A2"
ex_mov = [
    (datetime(2026, 10, 20, 18, 31), "#0001", "Registro", 1, "Diego", "Formulario"),
    (datetime(2026, 11, 1, 19, 5), "#0001", "Pedido", 1, "Diego", "Ticket 0012"),
    (datetime(2026, 11, 1, 19, 5), "#0001", "Hamburguesa adicional en el pedido", 1, "Diego", "Ticket 0012"),
    (datetime(2026, 11, 1, 19, 5), "#0001", "Bonus Montañera", 1, "Diego", "Ticket 0012"),
]
for r in range(2, ML + 1):
    if r - 2 < len(ex_mov):
        d, pid, acc, q, who, note = ex_mov[r - 2]
        for c, v in zip((1, 2, 4, 5, 7, 8), (d, pid, acc, q, who, note)):
            mv.cell(row=r, column=c, value=v)
        for c in (1, 2, 4, 5, 7, 8):
            mv.cell(row=r, column=c).fill = EX_FILL
    for c in (1, 2, 4, 5, 7, 8):
        mv.cell(row=r, column=c).font = IN_FONT
    mv.cell(row=r, column=1).number_format = "yyyy-mm-dd hh:mm"
    c3 = mv.cell(row=r, column=3, value=f'=IF(B{r}="","",IFERROR(INDEX(Jugadores!$C$2:$C${JL},MATCH(B{r},Jugadores!$A$2:$A${JL},0)),"⚠ ID no existe"))')
    c6 = mv.cell(row=r, column=6, value=f'=IF(OR(B{r}="",D{r}=""),"",IFERROR(INDEX({XP_RANGE},MATCH(D{r},{ACC_RANGE},0))*IF(E{r}="",1,E{r}),0))')
    for cell in (c3, c6):
        cell.font = CALC_FONT
        cell.fill = CALC_FILL
    c6.number_format = "#,##0"
dv_acc = DataValidation(type="list", formula1=ACC_RANGE, allow_blank=True)
mv.add_data_validation(dv_acc)
dv_acc.add(f"D2:D{ML}")
mv.conditional_formatting.add(f"C2:C{ML}", FormulaRule(formula=['LEFT(C2,1)="⚠"'], fill=PatternFill("solid", fgColor="F8C9C9")))

# ---------------- Ranking ----------------
rk = wb.create_sheet("Ranking")
rk["A1"] = "RANKING · MISIÓN PIXEL"
rk["A1"].font = TITLE
rk["A2"] = "Se actualiza solo. Para publicar en Historias: captura solo esta tabla (alias y XP, sin datos personales)."
rk["A2"].font = Font(name=F, italic=True, size=9)
header(rk, 4, ["Puesto", "Alias", "XP", "Premio si cerrara hoy"], [9, 24, 12, 34])
premios = {1: "GTA VI", 2: "2 hamburguesas a elección", 3: "2 hamburguesas a elección"}
for i in range(1, 21):
    r = 4 + i
    rk.cell(row=r, column=1, value=i).font = BASE
    rk.cell(row=r, column=2, value=f'=IFERROR(INDEX(Jugadores!$C$2:$C${JL},MATCH(A{r},Jugadores!$O$2:$O${JL},0)),"")').font = CALC_FONT
    c = rk.cell(row=r, column=3, value=f'=IFERROR(INDEX(Jugadores!$L$2:$L${JL},MATCH(A{r},Jugadores!$O$2:$O${JL},0)),"")')
    c.font = CALC_FONT
    c.number_format = "#,##0"
    prem = premios.get(i, "1 hamburguesa a elección" if i <= 10 else "")
    rk.cell(row=r, column=4, value=f'=IF(B{r}="","","{prem}")' if prem else "").font = BASE

# Resumen
rk["F4"] = "RESUMEN"
rk["F4"].font = SUB
resumen = [
    ("Jugadores registrados", f'=COUNTA(Jugadores!$A$2:$A${JL})'),
    ("Jugadores que compiten", f'=COUNTIF(Jugadores!$N$2:$N${JL},"Sí")'),
    ("XP total entregada", f'=SUM(Movimientos!$F$2:$F${ML})'),
    ("Pedidos con XP", f'=SUMIF(Movimientos!$D$2:$D${ML},"Pedido",Movimientos!$E$2:$E${ML})'),
    ("Hamburguesas registradas", f'=SUMIF(Movimientos!$D$2:$D${ML},"Pedido",Movimientos!$E$2:$E${ML})+SUMIF(Movimientos!$D$2:$D${ML},"Hamburguesa adicional en el pedido",Movimientos!$E$2:$E${ML})'),
    ("Referidos que compraron", f'=SUMIF(Movimientos!$D$2:$D${ML},"Referido con primer pedido",Movimientos!$E$2:$E${ML})'),
    ("Diseños de caja", f'=SUMIF(Movimientos!$D$2:$D${ML},"Diseño de caja #CajaPixel",Movimientos!$E$2:$E${ML})'),
    ("Recompensas por entregar", f'=COUNTIF(Jugadores!$R$2:$R${JL},"Sí*")'),
]
rk.column_dimensions["F"].width = 28
rk.column_dimensions["G"].width = 12
for i, (lab, fml) in enumerate(resumen, start=5):
    rk.cell(row=i, column=6, value=lab).font = BASE
    c = rk.cell(row=i, column=7, value=fml)
    c.font = CALC_FONT
    c.number_format = "#,##0"

# ---------------- Mensajes ----------------
ms = wb.create_sheet("Mensajes")
header(ms, 1, ["Momento", "Mensaje de WhatsApp (cambia lo que está entre [ ])"], [26, 110])
msgs = [
    ("Bienvenida (al registrarse)", "¡Misión aceptada, [nombre]! 🎮 Eres el Jugador [#ID] y ya tienes 100 XP. Tu power-up: salsa extra gratis en tu primer pedido. Di tu número de jugador cada vez que pidas para sumar XP. El 19 de noviembre el puesto 1 se lleva GTA VI. ¡Nos vemos en Pixel Burger!"),
    ("Después de un pedido", "+[XP] XP para el Jugador [#ID] 🍔 Ya vas en [total] XP. Te faltan [faltan] XP para tu próxima recompensa: [recompensa]."),
    ("Nivel desbloqueado", "🏆 ¡Logro desbloqueado! Llegaste a [nivel] y ganaste [recompensa]. Reclámala en tu próximo pedido diciendo tu número de jugador."),
    ("Ranking semanal (lunes, a todos)", "Ranking de la Misión Pixel 🕹️ Top 3: 1. [alias] · 2. [alias] · 3. [alias]. Quedan [días] días: el 19 de noviembre el puesto 1 se lleva GTA VI. Tú vas en el puesto [puesto] con [XP] XP."),
    ("Referido", "¿Tienes amigos con hambre? Que se registren con tu número de jugador [#ID]: cuando hagan su primer pedido, tú ganas 400 XP."),
    ("Cierre (18 de noviembre)", "⏳ Última noche de la Misión Pixel. Los pedidos hasta las 23:59 de hoy suman XP. Mañana anunciamos ganadores en el local."),
]
for r, (m, t) in enumerate(msgs, start=2):
    ms.cell(row=r, column=1, value=m).font = Font(name=F, bold=True, size=10)
    c = ms.cell(row=r, column=2, value=t)
    c.font = BASE
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ms.row_dimensions[r].height = 48

# ---------------- Formulario ----------------
fm = wb.create_sheet("Formulario")
header(fm, 1, ["Campo del formulario (Google Forms o Tally)", "Tipo", "Obligatorio"], [60, 26, 12])
campos = [
    ("Nombre completo", "Texto corto", "Sí"),
    ("Alias para el ranking (así aparecerás en público)", "Texto corto", "Sí"),
    ("Número de WhatsApp", "Texto corto (validar 10 dígitos)", "Sí"),
    ("Usuario de Instagram", "Texto corto", "No"),
    ("Barrio", "Lista: barrios del radio de domicilios + «Otro»", "Sí"),
    ("¿Te invitó un jugador? Escribe su número de jugador", "Texto corto", "No"),
    ("Confirmo que soy mayor de 18 años", "Casilla", "Sí"),
    ("Acepto las bases y condiciones de la Misión Pixel [enlace]", "Casilla", "Sí"),
    ("Autorizo a Pixel Burger a tratar mis datos personales para gestionar mi participación en la Misión Pixel, enviarme información del reto y promociones por WhatsApp, conforme a la Ley 1581 de 2012. Puedo consultar, actualizar o pedir la eliminación de mis datos escribiendo a [WhatsApp/correo].", "Casilla", "Sí"),
]
for r, row in enumerate(campos, start=2):
    for c, v in enumerate(row, start=1):
        cell = fm.cell(row=r, column=c, value=v)
        cell.font = BASE
        cell.alignment = Alignment(wrap_text=True, vertical="top")
fm.cell(row=len(campos) + 3, column=1, value="El texto de autorización de datos y las bases deben revisarse con un abogado antes de publicarse.").font = Font(name=F, italic=True, size=9)

wb.move_sheet("Ranking", offset=1 - wb.sheetnames.index("Ranking"))
wb.save(OUT)
print("ok", OUT, wb.sheetnames)
